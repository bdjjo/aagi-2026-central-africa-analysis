"""Structural audit of the supplied AAGI data-collection workbook.

Reports per-sheet dimensions, header rows, populated data rows, duplicates and
column missingness. Writes a JSON summary to data/processed/workbook_audit_raw.json.
Uses only the supplied workbook under data/raw/.
"""
from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WB_PATH = ROOT / "data" / "raw" / "CopyofAAGI_Data_Collection_centralafricarepublic.xlsx"
OUT = ROOT / "data" / "processed" / "workbook_audit_raw.json"

book = load_workbook(WB_PATH, data_only=False, read_only=True)
result = {"workbook": WB_PATH.name, "sheets": []}
for ws in book.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    max_cols = max((len(r) for r in rows), default=0)
    nonempty = [r for r in rows if any(v is not None and str(v).strip() != "" for v in r)]
    header_idx = next((i for i, r in enumerate(rows)
                       if any(v is not None and str(v).strip() != "" for v in r)), None)
    headers = list(rows[header_idx]) if header_idx is not None else []
    headers = [str(x).strip() if x is not None else "" for x in headers]
    data = rows[header_idx + 1:] if header_idx is not None else []
    frame = (pd.DataFrame(data, columns=headers + [f"unnamed_{i}" for i in range(len(headers), max_cols)])
             if headers else pd.DataFrame(data))
    frame = frame.dropna(how="all")
    dup_count = int(frame.duplicated().sum()) if not frame.empty else 0
    missing = {f"{i}:{c}": int(frame.iloc[:, i].isna().sum()) for i, c in enumerate(frame.columns)}
    result["sheets"].append({
        "name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column,
        "nonempty_rows": len(nonempty),
        "header_row_1indexed": (header_idx + 1 if header_idx is not None else None),
        "headers": headers, "data_rows_after_header": len(frame), "duplicate_rows": dup_count,
        "missing_by_column": missing,
    })

OUT.write_text(json.dumps(result, indent=2, default=str))
print("Wrote", OUT)
for s in result["sheets"]:
    print(s["name"], s["max_row"], s["max_column"], s["data_rows_after_header"])
